import io
import json
import os
import re
from datetime import date

import streamlit as st
from docx import Document
from openai import OpenAI, OpenAIError
from pypdf import PdfReader

st.set_page_config(
    page_title="AI Resume Reviewer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="collapsed",
)

MAX_FILE_MB = 5
MAX_RESUME_CHARS = 15_000   # ~4 pages; keeps input tokens bounded
MAX_JD_CHARS = 6_000
MAX_OUTPUT_TOKENS = 5_000   # typical review uses ~2.5k; hard ceiling on cost per call
MIN_TEXT_CHARS = 200
DEFAULT_MODEL = "gpt-5.4-mini"  # mini tier to keep costs low

SEVERITY_STYLE = {
    "high": ("#fde8e8", "#b42318", "High"),
    "medium": ("#fef4e6", "#b54708", "Medium"),
    "low": ("#eaf2ff", "#1d4ed8", "Low"),
}


# ---------------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------------

st.markdown(
    """
    <style>
    .block-container {padding-top: 2rem; padding-bottom: 3rem; max-width: 1150px;}
    #MainMenu, footer {visibility: hidden;}

    .hero {
        background: linear-gradient(135deg, #4338ca 0%, #6d28d9 55%, #9333ea 100%);
        border-radius: 18px; padding: 2rem 2.2rem; color: #fff; margin-bottom: 1.6rem;
    }
    .hero h1 {color: #fff; font-size: 2.1rem; margin: 0 0 .35rem 0; padding: 0;}
    .hero p {color: rgba(255,255,255,.88); font-size: 1.05rem; margin: 0;}
    .hero .pills {margin-top: 1rem; display: flex; gap: .5rem; flex-wrap: wrap;}
    .hero .pill {background: rgba(255,255,255,.16); border-radius: 999px;
                 padding: .25rem .8rem; font-size: .85rem;}

    .card {border: 1px solid rgba(128,128,128,.22); border-radius: 14px;
           padding: 1.1rem 1.25rem; margin-bottom: .9rem; background: rgba(128,128,128,.04);}
    .card h4 {margin: 0 0 .4rem 0; font-size: 1.02rem;}
    .card p {margin: .25rem 0; line-height: 1.5;}
    .muted {opacity: .72; font-size: .9rem;}

    .score-ring {display: flex; align-items: center; justify-content: center; flex-direction: column;
                 border-radius: 16px; padding: 1.2rem; color: #fff; text-align: center; height: 100%;}
    .score-ring .num {font-size: 3rem; font-weight: 800; line-height: 1;}
    .score-ring .lbl {font-size: .85rem; opacity: .9; margin-top: .35rem;}

    .badge {display: inline-block; border-radius: 999px; padding: .12rem .6rem;
            font-size: .75rem; font-weight: 600; margin-right: .4rem;}
    .chip {display: inline-block; border-radius: 8px; padding: .2rem .55rem; margin: .18rem .25rem .18rem 0;
           font-size: .85rem; border: 1px solid rgba(128,128,128,.3);}
    .chip.ok {background: #ecfdf3; color: #067647; border-color: #abefc6;}
    .chip.miss {background: #fef3f2; color: #b42318; border-color: #fecdca;}

    .before {border-left: 3px solid #f04438; padding: .35rem .8rem; margin: .4rem 0;
             background: rgba(240,68,56,.06); border-radius: 0 8px 8px 0;}
    .after {border-left: 3px solid #12b76a; padding: .35rem .8rem; margin: .4rem 0;
            background: rgba(18,183,106,.07); border-radius: 0 8px 8px 0;}

    div[data-testid="stFileUploader"] section {border-radius: 12px;}
    .stTabs [data-baseweb="tab-list"] {gap: .4rem;}
    .stTabs [data-baseweb="tab"] {padding: .5rem .9rem;}
    </style>
    """,
    unsafe_allow_html=True,
)


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_text(uploaded_file):
    """Return plain text from a PDF, DOCX or TXT upload."""
    data = uploaded_file.getvalue()
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected. Please upload an unlocked copy.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def clean_text(text):
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------
# LLM analysis
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """You are a senior recruiter and hiring manager with 15+ years of experience \
screening candidates across tech, product, business, design, finance and operations roles. \
You review resumes the way a real hiring panel does: evidence first, specific, candid, and \
always tied to the role the person is targeting.

STEP 1 - Gatekeeping. Decide whether the document is a resume/CV (a document describing one \
person's professional background: experience, education, skills). Cover letters, job \
descriptions, invoices, articles, assignments, reports, forms, etc. are NOT resumes. If it is \
not a resume, set "is_resume": false, fill "document_type" and "not_resume_reason", and leave \
every other field empty/zero.

STEP 2 - If it is a resume, analyse it against the TARGET ROLE. Use the role the user supplied; \
otherwise infer the most likely target role from the resume's trajectory. If a job description \
is supplied, treat it as the primary benchmark.

Rules for quality:
- Every gap and recommendation must be specific to THIS resume and THIS role. Quote or reference \
the actual resume content. Never give generic advice like "use action verbs" or "tailor your \
resume" without saying exactly where and how.
- Gaps = what a hiring manager for the target role expects to see but cannot find or finds weak: \
missing core skills/tools, lack of quantified impact, scope/seniority not demonstrated, missing \
domain experience, unexplained employment gaps, missing sections, ATS/formatting problems.
- Each fix must be an action the candidate can do this week (e.g. "Add a bullet under <Company> \
quantifying <X>: e.g. 'Cut onboarding time 30% by ...'"). Where the resume lacks the underlying \
experience, suggest how to acquire proof (project, certification, open-source, volunteer work).
- Bullet rewrites must reuse the candidate's real facts; use placeholders like [X%] where a \
number is needed but unknown. Never invent employers, titles, or achievements.
- Name the specific missing competencies for the target role and seniority (e.g. for a Senior \
PM: owning a P&L/metric, roadmap strategy, experimentation, leading without authority, \
stakeholder management). Say which job on the resume could plausibly carry the evidence.
- Employment gaps: list every role with its dates in your head, sort them, and flag any gap of \
roughly 6+ months between one role's end and the next role's start, or between the latest \
role's end and today. Today's date is {today}.
- Scores are honest, calibrated integers. 85+ is genuinely interview-ready for the role; \
most resumes land 55-75.

Return ONLY a JSON object with exactly this shape:
{{
  "is_resume": true,
  "document_type": "Resume",
  "not_resume_reason": "",
  "candidate": {{
    "name": "", "current_title": "", "seniority": "Entry | Mid | Senior | Lead/Manager | Executive",
    "years_experience": 0, "target_role": "", "target_role_inferred": true
  }},
  "overall_score": 0,
  "verdict": "one sentence a recruiter would say after a 30-second skim",
  "summary": "3-4 sentences: how competitive this resume is for the target role and why",
  "scores": {{"role_fit": 0, "impact": 0, "ats": 0, "clarity": 0}},
  "strengths": [{{"title": "", "evidence": "quote/reference from resume"}}],
  "gaps": [{{
    "area": "short label", "severity": "high | medium | low",
    "issue": "what is missing or weak, referencing the resume",
    "why_it_matters": "why a hiring manager for the target role cares",
    "fix": "concrete action with an example"
  }}],
  "bullet_rewrites": [{{"original": "", "improved": "", "why": ""}}],
  "keywords": {{"present": [""], "missing": [""]}},
  "jd_match": {{"match_percent": 0, "matched_requirements": [""], "missing_requirements": [""]}},
  "employment_gaps": [{{"period": "", "note": "how to address it on the resume / in interview"}}],
  "formatting_issues": [""],
  "action_plan": [{{"priority": 1, "action": "", "impact": "High | Medium", "effort": "15 min | 1 hour | 1 day | 1+ week"}}]
}}

Array sizes: strengths 3-5, gaps 4-8 (sorted by severity), bullet_rewrites 3-5 using the weakest \
real bullets, keywords present/missing 6-15 each (role-relevant terms an ATS would scan for), \
formatting_issues 0-5, action_plan exactly 5 sorted by priority. \
scores.* are 0-10; overall_score is 0-100. \
If no job description is given, return jd_match with match_percent 0 and empty arrays."""


def find_secret(name):
    """Look up a secret at the top level, inside any [section], or in the environment."""
    try:
        if st.secrets.get(name):
            return str(st.secrets[name]).strip()
        for value in st.secrets.values():
            if hasattr(value, "get") and value.get(name):
                return str(value[name]).strip()
    except Exception:
        pass
    return os.environ.get(name, "").strip() or None


def get_client():
    api_key = find_secret("OPENAI_API_KEY")
    if not api_key:
        return None
    return OpenAI(api_key=api_key, timeout=120, max_retries=2)


# Identical resume + role + JD within 24h is served from cache instead of re-billing the API.
@st.cache_data(ttl=86_400, max_entries=200, show_spinner=False)
def analyse_resume(_client, resume_text, target_role, job_description):
    user_parts = [f"TARGET ROLE: {target_role.strip() or 'Not provided - infer it from the resume'}"]
    if job_description.strip():
        user_parts.append(f"JOB DESCRIPTION:\n{job_description.strip()[:MAX_JD_CHARS]}")
    user_parts.append(f"DOCUMENT:\n{resume_text[:MAX_RESUME_CHARS]}")

    model = find_secret("OPENAI_MODEL") or DEFAULT_MODEL
    # GPT-5 family reasoning models only accept the default temperature.
    extra = {} if model.startswith(("gpt-5", "o")) else {"temperature": 0.3}
    response = _client.chat.completions.create(
        model=model,
        max_completion_tokens=MAX_OUTPUT_TOKENS,
        response_format={"type": "json_object"},
        **extra,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT.format(today=date.today().strftime("%B %Y"))},
            {"role": "user", "content": "\n\n".join(user_parts)},
        ],
    )
    return json.loads(response.choices[0].message.content)


# ---------------------------------------------------------------------------
# Rendering helpers
# ---------------------------------------------------------------------------

def esc(value):
    """HTML-escape model output before injecting it into markdown blocks."""
    return (
        str(value or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def score_color(score, scale=100):
    pct = score / scale
    if pct >= 0.8:
        return "linear-gradient(135deg,#059669,#10b981)"
    if pct >= 0.6:
        return "linear-gradient(135deg,#d97706,#f59e0b)"
    return "linear-gradient(135deg,#dc2626,#f43f5e)"


def as_int(value, default=0):
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return default


def chips(items, kind):
    items = [i for i in items or [] if str(i).strip()]
    if not items:
        return "<span class='muted'>None identified</span>"
    return "".join(f"<span class='chip {kind}'>{esc(i)}</span>" for i in items)


def build_report(result):
    """Plain-markdown version of the review for download."""
    c = result.get("candidate", {})
    lines = [
        f"# Resume Review - {c.get('name') or 'Candidate'}",
        f"**Target role:** {c.get('target_role', '')}  ",
        f"**Overall score:** {as_int(result.get('overall_score'))}/100  ",
        f"**Verdict:** {result.get('verdict', '')}",
        "",
        result.get("summary", ""),
        "",
        "## Action plan",
    ]
    for a in result.get("action_plan", []):
        lines.append(f"{a.get('priority')}. {a.get('action')} _(impact: {a.get('impact')}, effort: {a.get('effort')})_")
    lines += ["", "## Gaps"]
    for g in result.get("gaps", []):
        lines += [
            f"### {g.get('area')} ({g.get('severity', '').title()})",
            f"- **Issue:** {g.get('issue')}",
            f"- **Why it matters:** {g.get('why_it_matters')}",
            f"- **Fix:** {g.get('fix')}",
            "",
        ]
    lines.append("## Bullet rewrites")
    for b in result.get("bullet_rewrites", []):
        lines += [f"- **Before:** {b.get('original')}", f"  **After:** {b.get('improved')}", f"  _{b.get('why')}_", ""]
    kw = result.get("keywords", {})
    lines += ["## Keywords", f"**Missing:** {', '.join(kw.get('missing', []))}", f"**Present:** {', '.join(kw.get('present', []))}"]
    if result.get("employment_gaps"):
        lines += ["", "## Employment gaps"]
        lines += [f"- **{e.get('period')}:** {e.get('note')}" for e in result["employment_gaps"]]
    if result.get("formatting_issues"):
        lines += ["", "## Formatting / ATS"]
        lines += [f"- {f}" for f in result["formatting_issues"]]
    lines += ["", "## Strengths"]
    lines += [f"- **{s.get('title')}:** {s.get('evidence')}" for s in result.get("strengths", [])]
    return "\n".join(lines)


def render_not_resume(result):
    doc_type = result.get("document_type") or "a different kind of document"
    reason = result.get("not_resume_reason") or ""
    st.markdown(
        f"""
        <div class="card" style="border-color:#fdb022;background:rgba(253,176,34,.08);">
          <h4>🤔 This doesn't look like a resume</h4>
          <p>It reads like <b>{esc(doc_type)}</b>. {esc(reason)}</p>
          <p class="muted">Upload your CV / resume (the document listing your experience, education
          and skills) and we'll review it against your target role.</p>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_results(result, has_jd):
    c = result.get("candidate", {}) or {}
    overall = as_int(result.get("overall_score"))
    scores = result.get("scores", {}) or {}

    role = c.get("target_role") or "your target role"
    role_note = " (inferred)" if c.get("target_role_inferred") else ""
    meta = " · ".join(
        x for x in [c.get("current_title"), c.get("seniority"),
                    f"{c.get('years_experience')} yrs exp" if c.get("years_experience") else ""] if x
    )

    st.markdown(f"### Review for {esc(c.get('name') or 'your resume')}")
    st.markdown(f"<span class='muted'>{esc(meta)}</span>", unsafe_allow_html=True)

    col_score, col_summary = st.columns([1, 3], gap="medium")
    with col_score:
        st.markdown(
            f"""<div class="score-ring" style="background:{score_color(overall)}">
                <div class="num">{overall}</div><div class="lbl">Overall / 100</div></div>""",
            unsafe_allow_html=True,
        )
    with col_summary:
        st.markdown(
            f"""<div class="card"><h4>🎯 Target: {esc(role)}{role_note}</h4>
                <p><b>{esc(result.get('verdict'))}</b></p>
                <p>{esc(result.get('summary'))}</p></div>""",
            unsafe_allow_html=True,
        )

    m = st.columns(5 if has_jd else 4)
    m[0].metric("Role fit", f"{as_int(scores.get('role_fit'))}/10")
    m[1].metric("Impact", f"{as_int(scores.get('impact'))}/10")
    m[2].metric("ATS readiness", f"{as_int(scores.get('ats'))}/10")
    m[3].metric("Clarity", f"{as_int(scores.get('clarity'))}/10")
    if has_jd:
        m[4].metric("JD match", f"{as_int((result.get('jd_match') or {}).get('match_percent'))}%")

    st.write("")
    tab_names = ["🚀 Action plan", "🔍 Gaps", "✍️ Bullet rewrites", "🔑 Keywords"]
    if has_jd:
        tab_names.append("📋 JD match")
    tab_names.append("💪 Strengths")
    tabs = st.tabs(tab_names)
    t = dict(zip(tab_names, tabs))

    with t["🚀 Action plan"]:
        st.caption("Do these in order - the top items move the needle most for this role.")
        for a in sorted(result.get("action_plan", []), key=lambda x: as_int(x.get("priority"), 99)):
            st.markdown(
                f"""<div class="card"><h4>{as_int(a.get('priority'))}. {esc(a.get('action'))}</h4>
                    <span class="badge" style="background:#ecfdf3;color:#067647">Impact: {esc(a.get('impact'))}</span>
                    <span class="badge" style="background:#eef4ff;color:#3538cd">Effort: {esc(a.get('effort'))}</span>
                    </div>""",
                unsafe_allow_html=True,
            )

    with t["🔍 Gaps"]:
        gaps = result.get("gaps", [])
        order = {"high": 0, "medium": 1, "low": 2}
        for g in sorted(gaps, key=lambda x: order.get(str(x.get("severity", "")).lower(), 3)):
            bg, fg, label = SEVERITY_STYLE.get(str(g.get("severity", "")).lower(), SEVERITY_STYLE["low"])
            st.markdown(
                f"""<div class="card">
                    <h4><span class="badge" style="background:{bg};color:{fg}">{label}</span>{esc(g.get('area'))}</h4>
                    <p><b>What's missing:</b> {esc(g.get('issue'))}</p>
                    <p><b>Why it matters:</b> {esc(g.get('why_it_matters'))}</p>
                    <p><b>✅ Fix:</b> {esc(g.get('fix'))}</p></div>""",
                unsafe_allow_html=True,
            )
        emp_gaps = result.get("employment_gaps") or []
        if emp_gaps:
            st.markdown("##### ⏳ Employment gaps")
            for e in emp_gaps:
                st.markdown(
                    f"<div class='card'><h4>{esc(e.get('period'))}</h4><p>{esc(e.get('note'))}</p></div>",
                    unsafe_allow_html=True,
                )
        fmt = [f for f in result.get("formatting_issues") or [] if str(f).strip()]
        if fmt:
            st.markdown("##### 🧾 Formatting & ATS issues")
            st.markdown("".join(f"- {esc(f)}\n" for f in fmt))

    with t["✍️ Bullet rewrites"]:
        st.caption("Your weakest bullets, rewritten. Replace anything in [brackets] with your real numbers.")
        for b in result.get("bullet_rewrites", []):
            st.markdown(
                f"""<div class="card">
                    <div class="before"><span class="muted">Before</span><br>{esc(b.get('original'))}</div>
                    <div class="after"><span class="muted">After</span><br>{esc(b.get('improved'))}</div>
                    <p class="muted">💡 {esc(b.get('why'))}</p></div>""",
                unsafe_allow_html=True,
            )

    with t["🔑 Keywords"]:
        kw = result.get("keywords", {}) or {}
        st.markdown(
            f"""<div class="card"><h4>Missing - add these where you genuinely have the experience</h4>
                {chips(kw.get('missing'), 'miss')}</div>
                <div class="card"><h4>Already present</h4>{chips(kw.get('present'), 'ok')}</div>""",
            unsafe_allow_html=True,
        )

    if has_jd:
        with t["📋 JD match"]:
            jd = result.get("jd_match", {}) or {}
            pct = as_int(jd.get("match_percent"))
            st.progress(min(max(pct, 0), 100) / 100, text=f"{pct}% of the job's requirements are evidenced on your resume")
            st.markdown(
                f"""<div class="card"><h4>❌ Requirements not evidenced</h4>{chips(jd.get('missing_requirements'), 'miss')}</div>
                    <div class="card"><h4>✅ Requirements you meet</h4>{chips(jd.get('matched_requirements'), 'ok')}</div>""",
                unsafe_allow_html=True,
            )

    with t["💪 Strengths"]:
        st.caption("Keep these - and make sure they're visible in the top third of page one.")
        for s in result.get("strengths", []):
            st.markdown(
                f"<div class='card'><h4>{esc(s.get('title'))}</h4><p class='muted'>{esc(s.get('evidence'))}</p></div>",
                unsafe_allow_html=True,
            )

    st.write("")
    safe_name = re.sub(r"[^A-Za-z0-9]+", "_", c.get("name") or "resume").strip("_") or "resume"
    st.download_button(
        "⬇️ Download full review (Markdown)",
        data=build_report(result),
        file_name=f"{safe_name}_review.md",
        mime="text/markdown",
        use_container_width=True,
    )


# ---------------------------------------------------------------------------
# Page
# ---------------------------------------------------------------------------

st.markdown(
    """
    <div class="hero">
      <h1>📄 AI Resume Reviewer</h1>
      <p>Get a recruiter-grade review of your resume - the gaps that matter for <i>your</i> target role,
      and exactly what to change.</p>
      <div class="pills">
        <span class="pill">🎯 Role-specific gap analysis</span>
        <span class="pill">✍️ Bullet rewrites</span>
        <span class="pill">🔑 ATS keywords</span>
        <span class="pill">📋 Job description match</span>
      </div>
    </div>
    """,
    unsafe_allow_html=True,
)

client = get_client()
if client is None:
    st.error("The reviewer isn't configured yet (missing `OPENAI_API_KEY` in the app's secrets). Please check back soon.")
    st.stop()

with st.form("review_form", border=False):
    left, right = st.columns([1, 1], gap="large")
    with left:
        uploaded_file = st.file_uploader(
            "**1. Upload your resume**",
            type=["pdf", "docx", "txt"],
            help=f"PDF, DOCX or TXT, up to {MAX_FILE_MB} MB. Text-based files work best (not scanned images).",
        )
        target_role = st.text_input(
            "**2. Target role** (recommended)",
            placeholder="e.g. Senior Product Manager, Data Analyst, Frontend Engineer",
            help="Leave blank and we'll infer it from your resume.",
        )
    with right:
        job_description = st.text_area(
            "**3. Job description** (optional)",
            placeholder="Paste a job posting for a targeted match score and missing-requirement check...",
            height=205,
        )
    submitted = st.form_submit_button("🚀 Review my resume", type="primary", use_container_width=True)

if submitted:
    st.session_state.pop("result", None)
    if uploaded_file is None:
        st.warning("Please upload your resume first.")
        st.stop()
    if uploaded_file.size > MAX_FILE_MB * 1024 * 1024:
        st.warning(f"That file is larger than {MAX_FILE_MB} MB. Please upload a smaller copy.")
        st.stop()

    with st.status("Reviewing your resume...", expanded=True) as status:
        st.write("📖 Reading the document...")
        try:
            text = clean_text(extract_text(uploaded_file))
        except ValueError as e:
            status.update(label="Couldn't read that file", state="error")
            st.error(str(e))
            st.stop()
        except Exception:
            status.update(label="Couldn't read that file", state="error")
            st.error("We couldn't read that file - it may be corrupted. Try re-exporting it as a PDF or DOCX.")
            st.stop()

        if len(text) < MIN_TEXT_CHARS:
            status.update(label="Not enough text found", state="error")
            st.warning(
                "We couldn't find enough text in this file. If it's a scanned image or photo, "
                "please upload a text-based PDF or DOCX (e.g. export from Word / Google Docs)."
            )
            st.stop()

        st.write("🧠 Comparing it against what hiring managers expect for the role...")
        try:
            result = analyse_resume(client, text, target_role.strip(), job_description.strip())
        except json.JSONDecodeError:
            status.update(label="Something went wrong", state="error")
            st.error("The review came back in an unexpected format. Please try again.")
            st.stop()
        except OpenAIError:
            status.update(label="Something went wrong", state="error")
            st.error("The AI service is busy or unavailable right now. Please try again in a minute.")
            st.stop()

        status.update(label="Review ready!", state="complete", expanded=False)

    st.session_state["result"] = result
    st.session_state["has_jd"] = bool(job_description.strip())

if "result" in st.session_state:
    result = st.session_state["result"]
    st.divider()
    if not result.get("is_resume", True):
        render_not_resume(result)
    else:
        render_results(result, st.session_state.get("has_jd", False))
else:
    st.write("")
    c1, c2, c3 = st.columns(3)
    for col, (icon, title, body) in zip(
        (c1, c2, c3),
        (
            ("🔍", "Finds real gaps", "Missing skills, weak impact, unexplained gaps - judged against your target role, not a generic checklist."),
            ("✅", "Tells you the fix", "Every gap comes with a concrete change you can make this week, plus rewritten bullets."),
            ("🔒", "Private", "Your file is read in memory for this review and is not saved by this app."),
        ),
    ):
        col.markdown(f"<div class='card'><h4>{icon} {title}</h4><p class='muted'>{body}</p></div>", unsafe_allow_html=True)
